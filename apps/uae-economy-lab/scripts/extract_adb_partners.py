#!/usr/bin/env python3
"""Extract real UAE bilateral sector flows from ADB's public 2024 MRIO.

Requires openpyxl. Downloads about 100MB if --source is not already present.
Run: python scripts/extract_adb_partners.py --source /path/to/adb-mrio74-2024.xlsx
"""
import argparse
import hashlib
import json
from pathlib import Path
import urllib.request

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
URL = "https://kidb.adb.org/download/mrio-file/445"


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", type=Path, default=ROOT / "data/source/adb-mrio74-2024.xlsx")
    args = parser.parse_args()
    if not args.source.exists():
        args.source.parent.mkdir(parents=True, exist_ok=True)
        with urllib.request.urlopen(URL, timeout=120) as response, args.source.open("wb") as target:
            while chunk := response.read(1024 * 1024):
                target.write(chunk)
    baseline_path = ROOT / "src/data/uae-baseline.json"
    baseline = json.loads(baseline_path.read_text())
    workbook = openpyxl.load_workbook(args.source, read_only=True, data_only=True)
    names = {r[0]: r[1] for r in workbook["Legend"].iter_rows(min_row=2, values_only=True) if r[0]}
    rows = workbook["ADB MRIO 2024"].iter_rows(values_only=True)
    header = [next(rows) for _ in range(7)]
    countries, industries = header[5], header[6]
    columns = {country: [j for j, v in enumerate(countries) if v == country and industries[j]] for country in names}
    assert "UAE" in columns and len(columns["UAE"]) == 40
    partners = {country: {"id": country, "name": names[country], "importsBySector": [0.] * 35, "exportsBySector": [0.] * 35} for country in names if country != "UAE"}
    uaecols = columns["UAE"]
    domestic = [[0.] * 35 for _ in range(35)]
    for index, row in enumerate(rows, start=8):
        country, sector = row[2], row[3]
        if country not in names or not isinstance(sector, str) or not sector.startswith("c"):
            continue
        i = int(sector[1:]) - 1
        if country == "UAE":
            domestic[i] = [float(row[j] or 0) for j in uaecols if industries[j].startswith("c")]
            for code, partner in partners.items():
                partner["exportsBySector"][i] = sum(float(row[j] or 0) for j in columns[code])
        else:
            partners[country]["importsBySector"][i] = sum(float(row[j] or 0) for j in uaecols)
    partner_list = sorted(partners.values(), key=lambda p: -(sum(p["importsBySector"]) + sum(p["exportsBySector"])))
    export_residuals = [sum(p["exportsBySector"][i] for p in partner_list) - s["exports"] for i, s in enumerate(baseline["sectors"])]
    import_residuals = [sum(p["importsBySector"][i] for p in partner_list) - s["imports"] for i, s in enumerate(baseline["sectors"])]
    domestic_residual = max(abs(domestic[i][j] - baseline["domesticIO"][i][j]) for i in range(35) for j in range(35))
    validation = {"maxExportResidual": max(map(abs, export_residuals)), "maxImportResidual": max(map(abs, import_residuals)), "maxDomesticIOResidual": domestic_residual, "toleranceUSDmillion": .00001}
    validation["balanced"] = all(validation[k] <= .00001 for k in ["maxExportResidual", "maxImportResidual", "maxDomesticIOResidual"])
    result = {
        "meta": {"source": "ADB 74-economy MRIO plus rest of world, 2024", "year": 2024, "units": "USD million, current prices", "edition": "August 2025", "sourceUrl": URL, "sha256": hashlib.sha256(args.source.read_bytes()).hexdigest(), "retrievedAt": "2026-09-14", "definition": "Exports by UAE supplying industry to each foreign economy's industries and five final uses. Imports by foreign supplying industry into UAE industries and five final uses. Values summed directly without rescaling.", "limitations": ["Partner coverage follows ADB MRIO. Countries not individually represented are included in rest of world.", "Values are MRIO flows and are not interchangeable with customs commodity trade totals."]},
        "partners": partner_list, "validation": validation,
    }
    target = ROOT / "src/data/uae-partners.json"
    target.write_text(json.dumps(result, indent=2, allow_nan=False) + "\n")
    if not validation["balanced"]:
        raise ValueError(f"MRIO and national source do not agree; partner data left separate: {validation}")
    baseline["partners"] = partner_list
    baseline["meta"]["partners"] = result["meta"]
    baseline["validation"]["partners"] = validation
    baseline["meta"]["limitations"] = [v for v in baseline["meta"]["limitations"] if not v.startswith("The national table supplies")]
    baseline_path.write_text(json.dumps(baseline, indent=2, allow_nan=False) + "\n")
    print(json.dumps({"file": str(target), "partners": len(partner_list), "validation": validation, "top5": [{"name": p["name"], "imports": sum(p["importsBySector"]), "exports": sum(p["exportsBySector"])} for p in partner_list[:5]]}, indent=2))


if __name__ == "__main__":
    main()
